#!/usr/bin/env python3
"""Generate the M08 OFFER evaluation workbook from governed doctrine/control surfaces.

The workbook is a downstream review surface only. It never promotes compliance,
engineering, acceptance or award authority. Missing static BT weights and bidder
inputs remain blank; the generator does not invent them.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

EXPECTED_OFFERS = [f"OFFER-{i:02d}" for i in range(1, 51)]
EXPECTED_CLUSTERS = [f"C{i}" for i in range(1, 9)]
SHEETS = ["CONFIG", "STATIC_BT", "INPUT_AB", "CALC_AB", "CATEGORY_COMPARE", "DASHBOARD"]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_yaml(path: Path) -> dict[str, Any]:
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    if not isinstance(data, dict):
        raise AssertionError(f"{path} root must be mapping")
    return data


def offer_map(doctrine: dict[str, Any]) -> dict[str, dict[str, Any]]:
    clusters = doctrine["cluster_model"]["clusters"]
    weights = doctrine["cluster_weights"]
    out: dict[str, dict[str, Any]] = {}
    for cid in EXPECTED_CLUSTERS:
        offers = list(clusters[cid]["primary_offers"])
        for offer in offers:
            if offer in out:
                raise AssertionError(f"duplicate offer {offer}")
            out[offer] = {
                "cluster": cid,
                "cluster_name": clusters[cid]["name"],
                "cluster_weight": float(weights[cid]),
                "cluster_size": len(offers),
                "local_weight_default": 1.0 / len(offers),
            }
    if sorted(out) != EXPECTED_OFFERS:
        raise AssertionError("doctrine does not cover OFFER-01..50 exactly")
    return out


def autosize(ws) -> None:
    for col in range(1, ws.max_column + 1):
        values = [str(ws.cell(r, col).value or "") for r in range(1, ws.max_row + 1)]
        ws.column_dimensions[get_column_letter(col)].width = min(max(10, max(map(len, values), default=10) + 2), 42)
    ws.freeze_panes = "A2"


def build(doctrine_path: Path, control_path: Path, out: Path) -> dict[str, Any]:
    doctrine = load_yaml(doctrine_path)
    control = load_yaml(control_path)
    omap = offer_map(doctrine)
    doctrine_sha = sha256(doctrine_path)
    control_sha = sha256(control_path)

    risk = control["risk_implementation"]
    if risk["selected_policy"] != "OPTION_A_SYSTEM_RISK_APPLIED_ONCE_AT_ITEM_LEVEL":
        raise AssertionError("MIP2A requires Option-A risk-once policy")
    if control["weighting_alignment"]["canonical_field"] != "BT_Weight":
        raise AssertionError("BT_Weight must remain the single canonical field")
    if control["pairwise_alignment"]["governing_S_domain"] != [-1, 0, 1]:
        raise AssertionError("pairwise sign domain drift")

    wb = Workbook()
    wb.remove(wb.active)
    for name in SHEETS:
        wb.create_sheet(name)

    cfg = wb["CONFIG"]
    cfg.append(["Parameter", "Value", "Authority/role"])
    cfg.append(["Max_risk_factor", float(risk["observed_config"]["value"]), "risk cap; Option-A once at item level"])
    cfg.append(["Doctrine_SHA256", doctrine_sha, "semantic input identity"])
    cfg.append(["Workbook_Control_SHA256", control_sha, "workbook implementation control identity"])
    cfg.append(["Generated_Output_Authority", "DOWNSTREAM_VIEW_ONLY", "zero formal credit"])

    st = wb["STATIC_BT"]
    st.append(["OFFER", "Cluster", "Cluster_Name", "Cluster_Weight", "Local_Weight_Default", "BT_Weight", "BT_Weight_State"])
    for offer in EXPECTED_OFFERS:
        m = omap[offer]
        st.append([offer, m["cluster"], m["cluster_name"], m["cluster_weight"], m["local_weight_default"], None, "PENDING_SOURCE_OR_REVIEW"])

    inp = wb["INPUT_AB"]
    inp.append(["OFFER", "RawScore_A", "Confidence_A", "BaseRisk_A", "BidderRisk_A", "SystemRisk_A", "RawScore_B", "Confidence_B", "BaseRisk_B", "BidderRisk_B", "SystemRisk_B", "Reviewer_Notes"])
    for offer in EXPECTED_OFFERS:
        inp.append([offer] + [None] * 11)

    calc = wb["CALC_AB"]
    calc.append(["OFFER", "BT_Weight", "Pairwise_S", "BT_Contribution", "TotalRisk_A", "Adjusted_A", "TotalRisk_B", "Adjusted_B", "WeightedDelta_Diagnostic"])
    for r, offer in enumerate(EXPECTED_OFFERS, start=2):
        calc.append([
            offer,
            f"=STATIC_BT!F{r}",
            f'=IF(OR(INPUT_AB!B{r}="",INPUT_AB!G{r}=""),"",SIGN(INPUT_AB!B{r}-INPUT_AB!G{r}))',
            f'=IF(OR(B{r}="",C{r}=""),"",B{r}*C{r})',
            f'=IF(COUNTA(INPUT_AB!D{r}:F{r})=0,"",MIN(CONFIG!$B$2,SUM(INPUT_AB!D{r}:F{r})))',
            f'=IF(OR(INPUT_AB!B{r}="",E{r}=""),"",INPUT_AB!B{r}*(1-E{r}))',
            f'=IF(COUNTA(INPUT_AB!I{r}:K{r})=0,"",MIN(CONFIG!$B$2,SUM(INPUT_AB!I{r}:K{r})))',
            f'=IF(OR(INPUT_AB!G{r}="",G{r}=""),"",INPUT_AB!G{r}*(1-G{r}))',
            f'=IF(OR(B{r}="",F{r}="",H{r}=""),"",B{r}*(F{r}-H{r}))',
        ])

    cat = wb["CATEGORY_COMPARE"]
    cat.append(["Cluster", "Cluster_Name", "Cluster_Weight", "BT_Sum", "Risk_Adjusted_A_Avg", "Risk_Adjusted_B_Avg", "Diagnostic_Delta"])
    clusters = doctrine["cluster_model"]["clusters"]
    for i, cid in enumerate(EXPECTED_CLUSTERS, start=2):
        offers = list(clusters[cid]["primary_offers"])
        rows = [EXPECTED_OFFERS.index(o) + 2 for o in offers]
        bt_terms = ",".join(f"CALC_AB!D{r}" for r in rows)
        a_terms = ",".join(f"CALC_AB!F{r}" for r in rows)
        b_terms = ",".join(f"CALC_AB!H{r}" for r in rows)
        cat.append([
            cid,
            clusters[cid]["name"],
            float(doctrine["cluster_weights"][cid]),
            f'=IF(COUNTA({bt_terms})=0,"",SUM({bt_terms}))',
            f'=IF(COUNTA({a_terms})=0,"",AVERAGE({a_terms}))',
            f'=IF(COUNTA({b_terms})=0,"",AVERAGE({b_terms}))',
            f'=IF(OR(E{i}="",F{i}=""),"",E{i}-F{i})',
        ])

    dash = wb["DASHBOARD"]
    dash.append(["Metric", "Value", "Interpretation"])
    dash.append(["Pairwise_BT_Total", '=IF(COUNTA(CALC_AB!D2:D51)=0,"",SUM(CALC_AB!D2:D51))', "sign-based fixed-weight pairwise; not award authority"])
    dash.append(["Risk_Adjusted_A_Avg", '=IF(COUNTA(CALC_AB!F2:F51)=0,"",AVERAGE(CALC_AB!F2:F51))', "risk overlay; separate from requirement importance"])
    dash.append(["Risk_Adjusted_B_Avg", '=IF(COUNTA(CALC_AB!H2:H51)=0,"",AVERAGE(CALC_AB!H2:H51))', "risk overlay; separate from requirement importance"])
    dash.append(["WeightedDelta_Diagnostic_Total", '=IF(COUNTA(CALC_AB!I2:I51)=0,"",SUM(CALC_AB!I2:I51))', "diagnostic only; MUST NOT be called Bradley-Terry"])
    dash.append(["Formal_Credit", 0, "engineering/compliance/negotiation/acceptance/release/award"])

    for ws in wb.worksheets:
        autosize(ws)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)

    reloaded = load_workbook(out, data_only=False)
    if reloaded.sheetnames != SHEETS:
        raise AssertionError(f"sheet drift {reloaded.sheetnames}")
    with zipfile.ZipFile(out) as zf:
        table_parts = [n for n in zf.namelist() if n.startswith("xl/tables/")]
    if table_parts:
        raise AssertionError(f"unexpected Excel table parts: {table_parts}")

    receipt = {
        "status": "PASS_OPENPYXL_ROUNDTRIP_NO_TABLE_PART",
        "output": out.name,
        "output_sha256": sha256(out),
        "doctrine_sha256": doctrine_sha,
        "workbook_control_sha256": control_sha,
        "offer_count": 50,
        "sheets": SHEETS,
        "table_parts": [],
        "excel_native_clean_roundtrip": "WITHHELD_REQUIRES_EXCEL_RUNTIME",
        "authority_transfer": False,
        "formal_credit_delta": 0,
    }
    return receipt


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--doctrine", type=Path, default=Path("controls/QPS_OFFER_EVAL_SCORING_DOCTRINE_v1.yaml"))
    ap.add_argument("--control", type=Path, default=Path("controls/QPS_OFFER_EVAL_WORKBOOK_DELTA_2026-09-14_v1.yaml"))
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument("--receipt", type=Path)
    args = ap.parse_args()
    receipt = build(args.doctrine, args.control, args.out)
    if args.receipt:
        args.receipt.parent.mkdir(parents=True, exist_ok=True)
        args.receipt.write_text(json.dumps(receipt, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps(receipt, sort_keys=True))


if __name__ == "__main__":
    main()
