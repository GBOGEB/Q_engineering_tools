from pathlib import Path
import importlib.util
import zipfile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "qps_w183_generate_offer_eval_workbook.py"
DOCTRINE = ROOT / "controls" / "QPS_OFFER_EVAL_SCORING_DOCTRINE_v1.yaml"
CONTROL = ROOT / "controls" / "QPS_OFFER_EVAL_WORKBOOK_DELTA_2026-09-14_v1.yaml"

spec = importlib.util.spec_from_file_location("mip2a", SCRIPT)
mod = importlib.util.module_from_spec(spec)
assert spec.loader
spec.loader.exec_module(mod)


def build_tmp(tmp_path):
    out = tmp_path / "offer_eval.xlsx"
    receipt = mod.build(DOCTRINE, CONTROL, out)
    return out, receipt


def test_exact_workbook_stack_and_50_offer_rows(tmp_path):
    out, receipt = build_tmp(tmp_path)
    wb = load_workbook(out, data_only=False)
    assert wb.sheetnames == mod.SHEETS
    assert wb["STATIC_BT"].max_row == 51
    assert [wb["STATIC_BT"].cell(r, 1).value for r in range(2, 52)] == mod.EXPECTED_OFFERS
    assert receipt["offer_count"] == 50


def test_bt_weight_single_authority_and_missing_not_invented(tmp_path):
    out, _ = build_tmp(tmp_path)
    wb = load_workbook(out, data_only=False)
    headers = [c.value for c in wb["STATIC_BT"][1]]
    assert headers.count("BT_Weight") == 1
    assert all(wb["STATIC_BT"].cell(r, 6).value is None for r in range(2, 52))
    assert all(wb["STATIC_BT"].cell(r, 7).value == "PENDING_SOURCE_OR_REVIEW" for r in range(2, 52))


def test_pairwise_and_risk_are_separate(tmp_path):
    out, _ = build_tmp(tmp_path)
    wb = load_workbook(out, data_only=False)
    pairwise = wb["CALC_AB"]["C2"].value
    risk_a = wb["CALC_AB"]["E2"].value
    diagnostic = wb["CALC_AB"]["I2"].value
    assert "SIGN(INPUT_AB!B2-INPUT_AB!G2)" in pairwise
    assert "MIN(CONFIG!$B$2" in risk_a
    assert "F2-H2" in diagnostic
    assert "TotalRisk" not in pairwise
    assert "CALC_AB!E2" not in pairwise and "CALC_AB!G2" not in pairwise


def test_no_excel_table_part_and_native_excel_gate_withheld(tmp_path):
    out, receipt = build_tmp(tmp_path)
    with zipfile.ZipFile(out) as zf:
        assert not [n for n in zf.namelist() if n.startswith("xl/tables/")]
    assert receipt["status"] == "PASS_OPENPYXL_ROUNDTRIP_NO_TABLE_PART"
    assert receipt["excel_native_clean_roundtrip"] == "WITHHELD_REQUIRES_EXCEL_RUNTIME"


def test_authority_guard(tmp_path):
    _, receipt = build_tmp(tmp_path)
    assert receipt["authority_transfer"] is False
    assert receipt["formal_credit_delta"] == 0
