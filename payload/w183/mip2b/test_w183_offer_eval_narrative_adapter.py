from pathlib import Path
import importlib.util

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "qps_w183_generate_offer_eval_narrative.py"
DOCTRINE = ROOT / "controls" / "QPS_OFFER_EVAL_SCORING_DOCTRINE_v1.yaml"

spec = importlib.util.spec_from_file_location("mip2b", SCRIPT)
mod = importlib.util.module_from_spec(spec)
assert spec.loader
spec.loader.exec_module(mod)


def build_tmp(tmp_path):
    out = tmp_path / "narrative.docx"
    receipt = mod.build(DOCTRINE, out)
    return out, receipt


def doc_text(path):
    d = Document(path)
    return "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(
        cell.text for t in d.tables for row in t.rows for cell in row.cells
    )


def test_required_sections_and_12_principles(tmp_path):
    out, receipt = build_tmp(tmp_path)
    text = doc_text(out)
    for key in mod.EXPECTED_SECTIONS:
        assert mod.SECTION_TITLES[key] in text
    assert receipt["required_section_count"] == 11
    for i in range(1, 13):
        assert f"V6-{i:02d}" in text


def test_cluster_coverage_and_no_invented_guidance(tmp_path):
    out, receipt = build_tmp(tmp_path)
    text = doc_text(out)
    for cid in [f"C{i}" for i in range(1, 9)]:
        assert f"{cid} —" in text
        assert f"NOT_SPECIFIED_IN_DOCTRINE:{cid}.reviewer_focus" in text
        assert f"NOT_SPECIFIED_IN_DOCTRINE:{cid}.what_good_looks_like" in text
        assert f"NOT_SPECIFIED_IN_DOCTRINE:{cid}.primary_scoring_trap" in text
    assert receipt["cluster_count"] == 8


def test_authority_boundary_and_method_order(tmp_path):
    out, receipt = build_tmp(tmp_path)
    text = doc_text(out)
    assert "Static importance precedes bidder comparison" in text
    assert "Risk is represented as a separate delivery-confidence adjustment" in text
    assert "DOWNSTREAM_VIEW_ONLY" in text
    assert receipt["authority_transfer"] is False
    assert receipt["formal_credit_delta"] == 0
