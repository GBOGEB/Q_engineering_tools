#!/usr/bin/env python3
from __future__ import annotations
import argparse, hashlib, json
from pathlib import Path
from typing import Any
import yaml
from docx import Document
from docx.shared import Inches, Pt
EXPECTED_CLUSTERS=[f"C{i}" for i in range(1,9)]
FIXED_BLOCKS=["score_scale","confidence_scale","scoring_interpretation_note","bidder_A_score_confidence_notes","bidder_B_score_confidence_notes"]
VARIABLE_BLOCKS=["cluster_scope","primary_offer_list","reviewer_focus","what_good_looks_like","primary_scoring_trap","cross_cluster_dependency_cues"]
def sha256(path:Path)->str:return hashlib.sha256(path.read_bytes()).hexdigest()
def load_doctrine(path:Path)->dict[str,Any]:
 d=yaml.safe_load(path.read_text(encoding='utf-8')) or {}; c=d['one_pager_generation_contract']; assert c['one_page_per_cluster'] is True; assert c['fixed_blocks']==FIXED_BLOCKS; assert c['cluster_variable_blocks']==VARIABLE_BLOCKS; return d
def configure_compact_page(doc):
 s=doc.sections[0]; s.top_margin=Inches(.32); s.bottom_margin=Inches(.32); s.left_margin=Inches(.42); s.right_margin=Inches(.42)
 n=doc.styles['Normal']; n.font.name='Arial'; n.font.size=Pt(8); n.paragraph_format.space_after=Pt(1.5); n.paragraph_format.line_spacing=1.0
 t=doc.styles['Title']; t.font.name='Arial'; t.font.size=Pt(13); t.font.bold=True; t.paragraph_format.space_after=Pt(3)
def add_block(doc,title,body):
 p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1.5); p.paragraph_format.line_spacing=1.0; r=p.add_run(f'{title}: '); r.bold=True; r.font.size=Pt(8); b=p.add_run(body); b.font.size=Pt(8)
def fixed_payload(d):
 raw=d['model_boundary']['raw_score_scale']; conf=d['model_boundary']['confidence_scale']; return {'score_scale':'; '.join(f'{k}={v}' for k,v in raw.items()),'confidence_scale':'; '.join(f'{k}={v}' for k,v in conf.items()),'scoring_interpretation_note':'Reviewer score is not compliance status; confidence is not evidence maturity; gate failures remain non-compensating.','bidder_A_score_confidence_notes':'Reviewer input: Applicant A raw score, confidence, source/evidence note. No cross-bidder substitution.','bidder_B_score_confidence_notes':'Reviewer input: Applicant B raw score, confidence, source/evidence note. No cross-bidder substitution.'}
def build(doctrine_path:Path,out_dir:Path):
 d=load_doctrine(doctrine_path); dsha=sha256(doctrine_path); clusters=d['cluster_model']['clusters']; links=d.get('cross_cluster_links_no_weight_credit',[]); fixed=fixed_payload(d); out_dir.mkdir(parents=True,exist_ok=True); outputs=[]
 for cid in EXPECTED_CLUSTERS:
  info=clusters[cid]; doc=Document(); configure_compact_page(doc); doc.add_heading(f"Reviewer One-Pager — {cid} {info['name']}",level=0); meta=doc.add_paragraph(); meta.paragraph_format.space_after=Pt(2); m1=meta.add_run(f'Doctrine SHA-256: {dsha}\n'); m1.font.size=Pt(7.5); m2=meta.add_run('Classification: DOWNSTREAM_REVIEW_VIEW_ONLY | Formal credit: 0'); m2.font.size=Pt(7.5); m2.bold=True
  add_block(doc,'Score scale',fixed['score_scale']); add_block(doc,'Confidence scale',fixed['confidence_scale']); add_block(doc,'Scoring interpretation',fixed['scoring_interpretation_note']); add_block(doc,'Bidder A score / confidence / notes',fixed['bidder_A_score_confidence_notes']); add_block(doc,'Bidder B score / confidence / notes',fixed['bidder_B_score_confidence_notes']); add_block(doc,'Cluster scope',str(info['intent'])); add_block(doc,'Primary OFFER list',', '.join(info['primary_offers'])); add_block(doc,'Reviewer focus',f'NOT_SPECIFIED_IN_DOCTRINE:{cid}.reviewer_focus'); add_block(doc,'What good looks like',f'NOT_SPECIFIED_IN_DOCTRINE:{cid}.what_good_looks_like'); add_block(doc,'Primary scoring trap',f'NOT_SPECIFIED_IN_DOCTRINE:{cid}.primary_scoring_trap'); linked=[link for link in links if any(o in info['primary_offers'] for o in link)]; add_block(doc,'Cross-cluster dependency cues',str(linked or 'None declared')); p=doc.add_paragraph('Cross-cluster cues carry no duplicate weight credit.'); p.paragraph_format.space_before=Pt(1); p.paragraph_format.space_after=Pt(0); p.runs[0].font.size=Pt(7.5); p.runs[0].italic=True
  path=out_dir/f'OFFER_EVAL_REVIEWER_ONEPAGER_{cid}.docx'; doc.save(path); outputs.append({'cluster':cid,'file':path.name,'sha256':sha256(path),'primary_offers':list(info['primary_offers'])})
 all_offers=[o for x in outputs for o in x['primary_offers']]; assert len(outputs)==8 and len(all_offers)==50 and len(set(all_offers))==50
 return {'status':'PASS_EIGHT_SOURCE_DRIVEN_ONEPAGERS','doctrine_sha256':dsha,'output_count':8,'outputs':outputs,'fixed_blocks_identical':True,'not_specified_markers_preserved':True,'physical_one_page_contract':'PENDING_RENDER_COUNT_PROOF','authority_transfer':False,'formal_credit_delta':0}
def main():
 ap=argparse.ArgumentParser(); ap.add_argument('--doctrine',type=Path,required=True); ap.add_argument('--out-dir',type=Path,required=True); ap.add_argument('--receipt',type=Path); a=ap.parse_args(); r=build(a.doctrine,a.out_dir); 
 if a.receipt: a.receipt.parent.mkdir(parents=True,exist_ok=True); a.receipt.write_text(json.dumps(r,indent=2,sort_keys=True)+'\n',encoding='utf-8')
 print(json.dumps(r,sort_keys=True))
if __name__=='__main__': main()
