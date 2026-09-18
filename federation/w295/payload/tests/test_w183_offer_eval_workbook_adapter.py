from pathlib import Path
import importlib.util
import zipfile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "qps_w183_generate_offer_eval_workbook.py"
DOCTRINE = ROOT / "controls" / "QPS_OFFER_EVAL_SCORING_DOCTRINE_v1.yaml"
CONTROL = ROOT / "controls" / "QPS_OFFER_EVAL_WORKBOOK_DELTA_2026-09-14_v1.yaml"
EQUATION = ROOT / "controls" / "QPS_OFFER_EVAL_EQUATION_CURRENT_v1.yaml"

spec = importlib.util.spec_from_file_location("mip2a", SCRIPT)
mod = importlib.util.module_from_spec(spec)
assert spec.loader
spec.loader.exec_module(mod)


def build_tmp(tmp_path):
    out = tmp_path / "offer_eval.xlsx"
    receipt = mod.build(DOCTRINE, CONTROL, out, EQUATION)
    return out, receipt


def test_exact_workbook_stack_and_50_offer_rows(tmp_path):
    out, receipt = build_tmp(tmp_path)
    wb = load_workbook(out, data_only=False)
    assert wb.sheetnames == mod.SHEETS
    assert wb["STATIC_BT"].max_row == 51
    assert [wb["STATIC_BT"].cell(r, 1).value for r in range(2, 52)] == mod.EXPECTED_OFFERS
    assert receipt["offer_count"] == 50
    assert receipt["equation_version"] == "OFFER_EVAL_EQ_v1"


def test_bt_weight_single_authority_and_missing_not_invented(tmp_path):
    out, _ = build_tmp(tmp_path)
    wb = load_workbook(out, data_only=False)
    headers = [c.value for c in wb["STATIC_BT"][1]]
    assert headers.count("BT_Weight") == 1
    assert all(wb["STATIC_BT"].cell(r, 6).value is None for r in range(2, 52))
    assert all(wb["STATIC_BT"].cell(r, 7).value == "PENDING_SOURCE_OR_REVIEW" for r in range(2, 52))


def test_equation_v1_confidence_na_gate_and_risk_once(tmp_path):
    out, receipt = build_tmp(tmp_path)
    wb = load_workbook(out, data_only=False)
    input_headers = [c.value for c in wb["INPUT_AB"][1]]
    for field in ["Confidence_A", "Confidence_B", "Applicable_A", "Gate_A", "Applicable_B", "Gate_B"]:
        assert field in input_headers
    pairwise = wb["CALC_AB"]["C2"].value
    risk_a = wb["CALC_AB"]["E2"].value
    adjusted_a = wb["CALC_AB"]["F2"].value
    diagnostic = wb["CALC_AB"]["I2"].value
    conditioned_a = wb["CALC_AB"]["J2"].value
    renorm_a = wb["CALC_AB"]["R2"].value
    row_state_a = wb["CALC_AB"]["V2"].value
    cluster_score_a = wb["CATEGORY_COMPARE"]["E2"].value
    final_score_a = wb["DASHBOARD"]["B4"].value
    assert "SIGN(INPUT_AB!B2-INPUT_AB!G2)" in pairwise
    assert "MIN(CONFIG!$B$2" in risk_a
    assert "COUNT(INPUT_AB!D2:F2)<3" in risk_a
    assert "ISTEXT(INPUT_AB!D2)" in risk_a
    assert "COUNTA(INPUT_AB!D2:F2)" not in risk_a
    assert "INPUT_AB!B2*INPUT_AB!C2" in conditioned_a
    assert "J2*N2*(1-E2)" in adjusted_a
    assert "NOT(ISNUMBER(E2))" in adjusted_a
    assert "NOT(ISNUMBER(E2))" in row_state_a
    assert 'L2=0,"UNSCORABLE"' in cluster_score_a
    assert "COUNT(" in cluster_score_a and "<L2" in cluster_score_a
    assert 'COUNTIF(CATEGORY_COMPARE!L2:L9,0)>0,"UNSCORABLE"' in final_score_a
    assert 'COUNTIF(CALC_AB!V2:V51,"DEFER")>0,"DEFER"' in final_score_a
    assert 'COUNTIF(CALC_AB!V2:V51,"FAIL")>0,"FAIL"' in final_score_a
    assert "SUMIFS" in renorm_a
    assert "SUM($B$2:$B$51)" in diagnostic
    assert "TotalRisk" not in pairwise
    assert receipt["confidence_conditioning_implemented"] is True
    assert receipt["explicit_applicability_NA_implemented"] is True
    assert receipt["missing_defer_fail_preserved"] is True
    assert receipt["all_NA_cluster_unscorable"] is True
    assert receipt["risk_application_count"] == 1
    assert receipt["final_system_risk_subtraction"] is False


def test_no_excel_table_part_and_native_excel_gate_withheld(tmp_path):
    out, receipt = build_tmp(tmp_path)
    with zipfile.ZipFile(out) as zf:
        assert not [n for n in zf.namelist() if n.startswith("xl/tables/")]
    assert receipt["status"] == "PASS_OPENPYXL_EQUATION_V1_ROUNDTRIP_NO_TABLE_PART"
    assert receipt["excel_native_clean_roundtrip"] == "WITHHELD_REQUIRES_EXCEL_RUNTIME"


def test_authority_guard(tmp_path):
    _, receipt = build_tmp(tmp_path)
    assert receipt["authority_transfer"] is False
    assert receipt["formal_credit_delta"] == 0
