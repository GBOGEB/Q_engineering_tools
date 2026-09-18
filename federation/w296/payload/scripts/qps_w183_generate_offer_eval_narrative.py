#!/usr/bin/env python3
"""Generate the full V6 OFFER evaluation narrative from governed doctrine + equation control.

The DOCX is a downstream view. It does not invent missing cluster guidance and does
not create source, compliance, acceptance or award authority.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

import yaml
from docx import Document

EXPECTED_SECTIONS = [
    "governing_purpose",
    "recursion_and_idempotence",
    "workbook_execution_stack",
    "execution_sequence",
    "v6_twelve_principles",
    "bt_input_conditioning",
    "cluster_deep_dives",
    "one_pager_generation",
    "decision_synthesis",
    "audit_traceability",
    "engineering_handover",
]

SECTION_TITLES = {
    "governing_purpose": "Governing purpose",
    "recursion_and_idempotence": "Recursion and idempotence",
    "workbook_execution_stack": "Workbook execution stack",
    "execution_sequence": "Execution sequence",
    "v6_twelve_principles": "V6 twelve principles",
    "bt_input_conditioning": "BT input conditioning and canonical score/risk equations",
    "cluster_deep_dives": "Cluster deep dives",
    "one_pager_generation": "One-pager generation",
    "decision_synthesis": "Decision synthesis",
    "audit_traceability": "Audit traceability",
    "engineering_handover": "Engineering handover",
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_yaml(path: Path) -> dict[str, Any]:
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    if not isinstance(data, dict):
        raise AssertionError("YAML root must be mapping")
    return data


def load_doctrine(path: Path) -> dict[str, Any]:
    data = load_yaml(path)
    required = data["narrative_generation_contract"]["required_sections"]
    if required != EXPECTED_SECTIONS:
        raise AssertionError(f"required section drift: {required}")
    return data


def add_kv(doc: Document, key: str, value: Any) -> None:
    p = doc.add_paragraph()
    p.add_run(f"{key}: ").bold = True
    p.add_run(str(value))


def add_equation(doc: Document, label: str, expression: str) -> None:
    p = doc.add_paragraph()
    p.add_run(f"{label}: ").bold = True
    p.add_run(expression)


def not_specified(field: str) -> str:
    return f"NOT_SPECIFIED_IN_DOCTRINE:{field}"


def build(
    doctrine_path: Path,
    out: Path,
    equation_path: Path | None = None,
) -> dict[str, Any]:
    d = load_doctrine(doctrine_path)
    equation_path = equation_path or (doctrine_path.parent / "QPS_OFFER_EVAL_EQUATION_CURRENT_v1.yaml")
    eq = load_yaml(equation_path)
    if eq["equation_version"] != "OFFER_EVAL_EQ_v1":
        raise AssertionError("equation version drift")
    dsha = sha256(doctrine_path)
    eqsha = sha256(equation_path)

    doc = Document()
    doc.add_heading("OFFER Evaluation Method — Full Recursive V6 Narrative", level=0)
    add_kv(doc, "Doctrine SHA-256", dsha)
    add_kv(doc, "Equation version", eq["equation_version"])
    add_kv(doc, "Equation control SHA-256", eqsha)
    add_kv(doc, "Classification", "DOWNSTREAM_VIEW_ONLY")
    add_kv(doc, "Formal credit", 0)

    doc.add_heading(SECTION_TITLES["governing_purpose"], level=1)
    doc.add_paragraph(d["purpose"])
    doc.add_paragraph("This narrative preserves W80 OFFER/RTM/EVAL and W152 evidence/V&V authority. It does not become source, compliance, acceptance or award authority.")

    doc.add_heading(SECTION_TITLES["recursion_and_idempotence"], level=1)
    doc.add_paragraph("The generated narrative is reconstructed from governed doctrine and one equation-control identity. Re-running the generator against unchanged inputs shall preserve section order, cluster membership, equation semantics and authority guards.")
    for rule in d.get("stop_rules", []):
        doc.add_paragraph(str(rule), style="List Bullet")

    doc.add_heading(SECTION_TITLES["workbook_execution_stack"], level=1)
    doc.add_paragraph("STATIC_BT -> INPUT_AB -> CALC_AB -> CATEGORY_COMPARE -> DASHBOARD")
    doc.add_paragraph("Static importance precedes bidder comparison. Risk is represented as a separate delivery-confidence adjustment and must not rewrite requirement importance.")

    doc.add_heading(SECTION_TITLES["execution_sequence"], level=1)
    for s in [
        "1. Freeze the requirement-importance prior independently of bidders.",
        "2. Extract bidder A and bidder B evidence without cross-bidder substitution.",
        "3. Apply applicability, reviewer response score and confidence without converting missing or N/A to zero.",
        "4. Apply hard gates as non-compensating conditions.",
        "5. Apply BaseRisk + BidderRisk + SystemRisk exactly once at item level, capped by the configured risk maximum.",
        "6. Perform the governing sign-based pairwise comparison using the fixed static importance prior.",
        "7. Synthesize cluster and system review views without converting review support into acceptance authority.",
    ]:
        doc.add_paragraph(s)

    doc.add_heading(SECTION_TITLES["v6_twelve_principles"], level=1)
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "ID"
    table.rows[0].cells[1].text = "Principle"
    table.rows[0].cells[2].text = "Interpretation"
    for item in d["v6_doctrine"]:
        cells = table.add_row().cells
        cells[0].text = str(item["id"])
        cells[1].text = str(item["principle"])
        cells[2].text = str(item["interpretation"])

    doc.add_heading(SECTION_TITLES["bt_input_conditioning"], level=1)
    doc.add_paragraph(
        "The following equations are the canonical OFFER_EVAL_EQ_v1 family. They replace the ambiguous "
        "legacy blended final-score expressions for production use while preserving the frozen method: "
        "static requirement importance remains separate from bidder-response scoring, and risk is applied once."
    )
    ce = eq["canonical_equations"]
    add_equation(doc, "Conditioned evaluator score", ce["conditioned_evaluator"]["expression"])
    add_equation(doc, "Applicable local-weight renormalization", ce["applicability_renormalization"]["expression"])
    add_equation(doc, "Raw risk", ce["raw_risk"]["expression"])
    add_equation(doc, "Capped total risk", ce["total_risk"]["expression"])
    add_equation(doc, "Risk-adjusted item input", ce["adjusted_item"]["expression"])
    add_equation(doc, "Cluster score", ce["cluster_score"]["expression"])
    add_equation(doc, "Final review score", ce["final_review_score"]["expression"])
    add_equation(doc, "Governing pairwise review signal", ce["governing_pairwise"]["expression"])
    add_equation(doc, "Optional diagnostic delta", ce["diagnostic_delta"]["expression"])
    doc.add_paragraph(
        "System risk is already included inside R_i_b and therefore inside X_i_b. No second final system-risk "
        "subtraction or second risk multiplier is permitted. BT(A,B) remains a separate static-prior-weighted "
        "pairwise review signal and is not blended into FinalReviewScore_b."
    )
    doc.add_paragraph(
        "N/A items are excluded and applicable local weights are renormalized. Missing-but-applicable evidence "
        "remains MISSING/DEFER/FAIL and is not converted to zero or a neutral score. A mandatory gate failure "
        "withholds a comparable final review score rather than being compensated by stronger scores elsewhere."
    )
    add_kv(doc, "Retired legacy blend", "Final = 0.70 * Evaluator + 0.30 * BT - Risk")
    add_kv(doc, "Retirement reason", "layer conflation and ambiguous/double risk application")
    c = d["scoring_execution_contract"]
    add_kv(doc, "Local item weighting", c["local_item_weighting"]["rule"])
    add_kv(doc, "Doctrine adjusted-item conceptual form", c["adjusted_item_input"]["conceptual_form"])
    add_kv(doc, "Pairwise transform status", c["pairwise_transform"]["status"])

    doc.add_heading(SECTION_TITLES["cluster_deep_dives"], level=1)
    clusters = d["cluster_model"]["clusters"]
    weights = d["cluster_weights"]
    links = d.get("cross_cluster_links_no_weight_credit", [])
    for cid in sorted(clusters):
        info = clusters[cid]
        doc.add_heading(f"{cid} — {info['name']}", level=2)
        add_kv(doc, "Provisional cluster weight", weights[cid])
        add_kv(doc, "Intent", info["intent"])
        add_kv(doc, "Primary OFFER items", ", ".join(info["primary_offers"]))
        linked = [link for link in links if any(o in info["primary_offers"] for o in link)]
        add_kv(doc, "Cross-cluster cues (no duplicate weight credit)", linked or "None declared")
        add_kv(doc, "Reviewer focus", not_specified(f"{cid}.reviewer_focus"))
        add_kv(doc, "What good looks like", not_specified(f"{cid}.what_good_looks_like"))
        add_kv(doc, "Primary scoring trap", not_specified(f"{cid}.primary_scoring_trap"))

    doc.add_heading(SECTION_TITLES["one_pager_generation"], level=1)
    op = d["one_pager_generation_contract"]
    add_kv(doc, "One page per cluster", op["one_page_per_cluster"])
    add_kv(doc, "Fixed blocks", ", ".join(op["fixed_blocks"]))
    add_kv(doc, "Cluster-variable blocks", ", ".join(op["cluster_variable_blocks"]))
    doc.add_paragraph(op["rule"])

    doc.add_heading(SECTION_TITLES["decision_synthesis"], level=1)
    doc.add_paragraph("Decision synthesis retains item traceability, cluster interpretation, confidence visibility and system-level coherence. Gate failures remain non-compensating and generated outputs do not create formal procurement award authority.")
    doc.add_paragraph("The static OFFER importance prior, commercial due-diligence action ranking and QPS TRIAGE execution pressure remain three different surfaces and shall not be substituted for one another.")

    doc.add_heading(SECTION_TITLES["audit_traceability"], level=1)
    add_kv(doc, "Doctrine path", doctrine_path.as_posix())
    add_kv(doc, "Doctrine SHA-256", dsha)
    add_kv(doc, "Equation path", equation_path.as_posix())
    add_kv(doc, "Equation SHA-256", eqsha)
    add_kv(doc, "OFFER denominator", d["model_boundary"]["offer_denominator"])
    add_kv(doc, "Primary cluster rule", d["cluster_model"]["rule"])
    add_kv(doc, "Generated binary role", d["narrative_generation_contract"]["generated_binary_role"])

    doc.add_heading(SECTION_TITLES["engineering_handover"], level=1)
    doc.add_paragraph("Engineering handover remains a QPS child re-entry matter. This DOCX is a review-support projection and carries zero engineering/compliance/negotiation/SCK-acceptance/release/award credit.")

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    receipt = {
        "status": "PASS_SOURCE_DRIVEN_NARRATIVE_EQUATION_V1_VISIBLE",
        "output": out.name,
        "output_sha256": sha256(out),
        "doctrine_sha256": dsha,
        "equation_control_sha256": eqsha,
        "equation_version": eq["equation_version"],
        "required_sections": EXPECTED_SECTIONS,
        "required_section_count": len(EXPECTED_SECTIONS),
        "cluster_count": len(clusters),
        "equation_visibility": True,
        "risk_application_count": 1,
        "final_system_risk_subtraction": False,
        "not_specified_markers_preserved": True,
        "authority_transfer": False,
        "formal_credit_delta": 0,
    }
    return receipt


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--doctrine", type=Path, default=Path("controls/QPS_OFFER_EVAL_SCORING_DOCTRINE_v1.yaml"))
    ap.add_argument("--equation", type=Path, default=Path("controls/QPS_OFFER_EVAL_EQUATION_CURRENT_v1.yaml"))
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument("--receipt", type=Path)
    args = ap.parse_args()
    receipt = build(args.doctrine, args.out, args.equation)
    if args.receipt:
        args.receipt.parent.mkdir(parents=True, exist_ok=True)
        args.receipt.write_text(json.dumps(receipt, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps(receipt, sort_keys=True))


if __name__ == "__main__":
    main()
