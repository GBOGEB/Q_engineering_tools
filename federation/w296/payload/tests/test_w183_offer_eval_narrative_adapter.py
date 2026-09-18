from pathlib import Path
import importlib.util

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "qps_w183_generate_offer_eval_narrative.py"
DOCTRINE = ROOT / "controls" / "QPS_OFFER_EVAL_SCORING_DOCTRINE_v1.yaml"
EQUATION = ROOT / "controls" / "QPS_OFFER_EVAL_EQUATION_CURRENT_v1.yaml"

spec = importlib.util.spec_from_file_location("mip2b", SCRIPT)
mod = importlib.util.module_from_spec(spec)
assert spec.loader
spec.loader.exec_module(mod)


def build_tmp(tmp_path):
    out = tmp_path / "narrative.docx"
    receipt = mod.build(DOCTRINE, out, EQUATION)
    return out, receipt


def doc_text(path):
    d = Document(path)
    return "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(
        cell.text for t in d.tables for row in t.rows for cell in row.cells
    )


def test_required_sections_and_12_principles(tmp_path):
    out, receipt = build_tmp(tmp_path)
    text = doc_text(out)
    for key in mod.EXPECTED_SECTIONS:
        assert mod.SECTION_TITLES[key] in text
    assert receipt["required_section_count"] == 11
    for i in range(1, 13):
        assert f"V6-{i:02d}" in text


def test_cluster_coverage_and_no_invented_guidance(tmp_path):
    out, receipt = build_tmp(tmp_path)
    text = doc_text(out)
    for cid in [f"C{i}" for i in range(1, 9)]:
        assert f"{cid} —" in text
        assert f"NOT_SPECIFIED_IN_DOCTRINE:{cid}.reviewer_focus" in text
        assert f"NOT_SPECIFIED_IN_DOCTRINE:{cid}.what_good_looks_like" in text
        assert f"NOT_SPECIFIED_IN_DOCTRINE:{cid}.primary_scoring_trap" in text
    assert receipt["cluster_count"] == 8


def test_equation_v1_visible_and_authority_boundary(tmp_path):
    out, receipt = build_tmp(tmp_path)
    text = doc_text(out)
    assert "Static importance precedes bidder comparison" in text
    assert "OFFER_EVAL_EQ_v1" in text
    assert "E_i_b = q_i_b * C_i_b" in text
    assert "R_i_b = MIN(R_max, RawRisk_i_b)" in text
    assert "X_i_b = E_i_b * G_i_b * (1 - R_i_b)" in text
    assert "FinalReviewScore_b = SUM_c(alpha_c * ClusterScore_c_b)" in text
    assert "BT(A,B) = SUM_i(W_i * S_i(A,B))" in text
    assert "No second final system-risk subtraction" in text
    assert "DOWNSTREAM_VIEW_ONLY" in text
    assert receipt["equation_visibility"] is True
    assert receipt["risk_application_count"] == 1
    assert receipt["final_system_risk_subtraction"] is False
    assert receipt["authority_transfer"] is False
    assert receipt["formal_credit_delta"] == 0
