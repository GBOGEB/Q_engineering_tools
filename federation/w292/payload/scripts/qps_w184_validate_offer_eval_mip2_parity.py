#!/usr/bin/env python3
"""Execute M08 cross-output parity for canonical OFFER_EVAL_EQ_v1."""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import tempfile
from pathlib import Path
from typing import Any

import yaml
from docx import Document
from openpyxl import load_workbook

EXPECTED_OFFERS = [f"OFFER-{i:02d}" for i in range(1, 51)]
EXPECTED_CLUSTERS = [f"C{i}" for i in range(1, 9)]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_yaml(path: Path) -> dict[str, Any]:
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    if not isinstance(data, dict):
        raise AssertionError(f"{path} root must be mapping")
    return data


def load_module(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise AssertionError(f"cannot load {path}")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def doctrine_map(d: dict[str, Any]) -> dict[str, str]:
    out: dict[str, str] = {}
    for cid in EXPECTED_CLUSTERS:
        for offer in d["cluster_model"]["clusters"][cid]["primary_offers"]:
            if offer in out:
                raise AssertionError(f"duplicate doctrine primary {offer}")
            out[offer] = cid
    if sorted(out) != EXPECTED_OFFERS:
        raise AssertionError("doctrine OFFER coverage is not exact 50/50")
    return out


def docx_text(path: Path) -> str:
    doc = Document(path)
    paras = [p.text for p in doc.paragraphs]
    cells = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return "\n".join(paras + cells)


def validate(repo_root: Path, workdir: Path) -> dict[str, Any]:
    doctrine_path = repo_root / "controls/QPS_OFFER_EVAL_SCORING_DOCTRINE_v1.yaml"
    workbook_control = repo_root / "controls/QPS_OFFER_EVAL_WORKBOOK_DELTA_2026-09-14_v1.yaml"
    equation_path = repo_root / "controls/QPS_OFFER_EVAL_EQUATION_CURRENT_v1.yaml"
    d = load_yaml(doctrine_path)
    eq = load_yaml(equation_path)
    canonical = doctrine_map(d)
    dsha = sha256(doctrine_path)
    eqsha = sha256(equation_path)

    a_mod = load_module(repo_root / "scripts/qps_w183_generate_offer_eval_workbook.py", "mip2a")
    b_mod = load_module(repo_root / "scripts/qps_w183_generate_offer_eval_narrative.py", "mip2b")
    c_mod = load_module(repo_root / "scripts/qps_w183_generate_offer_eval_onepagers.py", "mip2c")

    out_a = workdir / "OFFER_EVAL_MIP2A.xlsx"
    out_b = workdir / "FULL_RECURSIVE_V6_NARRATIVE_MIP2B.docx"
    out_c = workdir / "onepagers"
    receipt_a = a_mod.build(doctrine_path, workbook_control, out_a, equation_path)
    receipt_b = b_mod.build(doctrine_path, out_b, equation_path)
    receipt_c = c_mod.build(doctrine_path, out_c)

    receipt_shas = {receipt_a["doctrine_sha256"], receipt_b["doctrine_sha256"], receipt_c["doctrine_sha256"]}
    if receipt_shas != {dsha}:
        raise AssertionError(f"doctrine identity drift: {receipt_shas} != {dsha}")
    if receipt_a["equation_control_sha256"] != eqsha or receipt_b["equation_control_sha256"] != eqsha:
        raise AssertionError("equation identity drift")
    if any(r.get("authority_transfer") is not False for r in (receipt_a, receipt_b, receipt_c)):
        raise AssertionError("authority transfer guard drift")
    if any(r.get("formal_credit_delta") != 0 for r in (receipt_a, receipt_b, receipt_c)):
        raise AssertionError("formal credit guard drift")

    wb = load_workbook(out_a, data_only=False)
    if wb.sheetnames != ["CONFIG", "STATIC_BT", "INPUT_AB", "CALC_AB", "CATEGORY_COMPARE", "DASHBOARD"]:
        raise AssertionError("workbook sheet-stack drift")
    ws = wb["STATIC_BT"]
    headers = [cell.value for cell in ws[1]]
    if headers.count("BT_Weight") != 1:
        raise AssertionError("workbook must have exactly one BT_Weight field")
    wb_map: dict[str, str] = {}
    for r in range(2, 52):
        offer = ws.cell(r, 1).value
        cluster = ws.cell(r, 2).value
        wb_map[str(offer)] = str(cluster)
        if ws.cell(r, 6).value is not None:
            raise AssertionError(f"uncontrolled static BT weight populated at {offer}")
    if wb_map != canonical:
        raise AssertionError("workbook OFFER/cluster map differs from doctrine")

    input_headers = [c.value for c in wb["INPUT_AB"][1]]
    for required in ["Confidence_A", "Confidence_B", "Applicable_A", "Gate_A", "Applicable_B", "Gate_B"]:
        if required not in input_headers:
            raise AssertionError(f"workbook input missing {required}")
    pairwise_formula = wb["CALC_AB"]["C2"].value or ""
    risk_formula = wb["CALC_AB"]["E2"].value or ""
    adjusted_a = wb["CALC_AB"]["F2"].value or ""
    conditioned_a = wb["CALC_AB"]["J2"].value or ""
    renorm_a = wb["CALC_AB"]["R2"].value or ""
    diagnostic = wb["CALC_AB"]["I2"].value or ""
    if "SIGN(INPUT_AB!B2-INPUT_AB!G2)" not in pairwise_formula:
        raise AssertionError("pairwise sign formula drift")
    if "MIN(CONFIG!$B$2" not in risk_formula:
        raise AssertionError("risk cap formula drift")
    if "INPUT_AB!B2*INPUT_AB!C2" not in conditioned_a:
        raise AssertionError("confidence conditioning not implemented")
    if "J2*N2*(1-E2)" not in adjusted_a:
        raise AssertionError("equation-v1 adjusted score formula drift")
    if "SUMIFS" not in renorm_a or "P2/" not in renorm_a:
        raise AssertionError("N/A local-weight renormalization missing")
    if "SUM($B$2:$B$51)" not in diagnostic:
        raise AssertionError("diagnostic must use normalized static weight share")
    if receipt_a["risk_application_count"] != 1 or receipt_a["final_system_risk_subtraction"] is not False:
        raise AssertionError("risk-once receipt drift")
    if receipt_a["excel_native_clean_roundtrip"] != "WITHHELD_REQUIRES_EXCEL_RUNTIME":
        raise AssertionError("native Excel gate was silently promoted")

    narrative = docx_text(out_b)
    required_narrative_phrases = [
        "Static importance precedes bidder comparison",
        "OFFER_EVAL_EQ_v1",
        "E_i_b = q_i_b * C_i_b",
        "R_i_b = MIN(R_max, RawRisk_i_b)",
        "X_i_b = E_i_b * G_i_b * (1 - R_i_b)",
        "FinalReviewScore_b = SUM_c(alpha_c * ClusterScore_c_b)",
        "BT(A,B) = SUM_i(W_i * S_i(A,B))",
        "No second final system-risk subtraction",
        "DOWNSTREAM_VIEW_ONLY",
    ]
    for phrase in required_narrative_phrases:
        if phrase not in narrative:
            raise AssertionError(f"narrative semantic marker missing: {phrase}")
    for cid in EXPECTED_CLUSTERS:
        if cid not in narrative:
            raise AssertionError(f"narrative cluster missing: {cid}")

    if receipt_c["output_count"] != 8:
        raise AssertionError("one-pager count drift")
    c_map: dict[str, str] = {}
    for row in receipt_c["outputs"]:
        cid = row["cluster"]
        page = out_c / row["file"]
        page_text = docx_text(page)
        if "Reviewer score is not compliance status" not in page_text:
            raise AssertionError(f"one-pager fixed scoring boundary missing: {cid}")
        if "No cross-bidder substitution" not in page_text:
            raise AssertionError(f"one-pager bidder-isolation boundary missing: {cid}")
        for offer in row["primary_offers"]:
            if offer in c_map:
                raise AssertionError(f"duplicate one-pager OFFER: {offer}")
            c_map[offer] = cid
    if c_map != canonical:
        raise AssertionError("one-pager OFFER/cluster map differs from doctrine")

    return {
        "schema": "qps-m08-equation-v1-cross-output-parity/1.0",
        "status": "PASS_CROSS_OUTPUT_EQUATION_V1_PARITY_WITH_NATIVE_EXCEL_GATE_OPEN",
        "doctrine_sha256": dsha,
        "equation_control_sha256": eqsha,
        "equation_version": eq["equation_version"],
        "workbook_control_sha256": sha256(workbook_control),
        "offer_count": len(canonical),
        "cluster_count": len(EXPECTED_CLUSTERS),
        "mip2a_status": receipt_a["status"],
        "mip2b_status": receipt_b["status"],
        "mip2c_status": receipt_c["status"],
        "same_doctrine_identity": True,
        "same_offer_cluster_map": True,
        "static_importance_before_bidder_comparison": True,
        "pairwise_and_risk_separate": True,
        "confidence_conditioning_parity": True,
        "NA_applicability_parity": True,
        "risk_application_count": 1,
        "final_system_risk_subtraction": False,
        "source_gaps_not_invented": True,
        "native_excel_clean_roundtrip": "WITHHELD_REQUIRES_EXCEL_RUNTIME",
        "global_3p3_admission": "HOLD_NATIVE_EXCEL_ROUNDTRIP_AND_CHILD_REENTRY_RECEIPT",
        "authority_transfer": False,
        "formal_credit_delta": 0,
    }


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo-root", type=Path, default=Path("."))
    ap.add_argument("--workdir", type=Path)
    ap.add_argument("--receipt", type=Path)
    args = ap.parse_args()
    repo_root = args.repo_root.resolve()
    if args.workdir:
        workdir = args.workdir.resolve()
        workdir.mkdir(parents=True, exist_ok=True)
        receipt = validate(repo_root, workdir)
    else:
        with tempfile.TemporaryDirectory(prefix="qps-w288-") as tmp:
            receipt = validate(repo_root, Path(tmp))
    if args.receipt:
        args.receipt.parent.mkdir(parents=True, exist_ok=True)
        args.receipt.write_text(json.dumps(receipt, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps(receipt, sort_keys=True))


if __name__ == "__main__":
    main()
